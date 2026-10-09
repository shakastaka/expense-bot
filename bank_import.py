"""Conservative import of header-based bank CSV and Excel statements."""
import csv
import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

ALIASES = {
    'day': ('date', 'booking date', 'transaction date', 'value date', 'datum', 'datum knjizenja', 'datum knjiženja', 'datum valute', 'дата', 'дата операции', 'дата выполнения'),
    'description': ('description', 'details', 'merchant', 'narrative', 'payment details', 'opis', 'opis transakcije', 'namen', 'namen placila', 'namen plačila', 'prejemnik', 'описание', 'назначение платежа', 'назначение', 'получатель'),
    'amount': ('amount', 'transaction amount', 'znesek', 'znesek transakcije', 'сумма', 'сумма операции'),
    'debit': ('debit', 'withdrawal', 'money out', 'paid out', 'odliv', 'breme', 'v breme', 'izdatki', 'расход', 'списание'),
    'credit': ('credit', 'deposit', 'money in', 'paid in', 'priliv', 'dobro', 'v dobro', 'prejemki', 'приход', 'поступление'),
    'currency': ('currency', 'currency code', 'ccy', 'valuta', 'валюта'),
    'direction': ('type', 'transaction type', 'direction', 'tip', 'vrsta', 'тип операции'),
    'status': ('status', 'state', 'stanje', 'статус'),
    'account': ('account', 'account number', 'iban', 'racun', 'račun', 'номер счета'),
    'balance': ('balance', 'saldo', 'stanje racuna', 'остаток', 'остаток средств'),
}
CURRENCIES = {'€': 'EUR', '$': 'USD', '£': 'GBP', '₽': 'RUB'}

def normal(value):
    value = ''.join(c for c in unicodedata.normalize('NFKD', str(value).casefold()) if not unicodedata.combining(c))
    return ' '.join(re.findall(r'\w+', value))

LOOKUP = {normal(alias): field for field, aliases in ALIASES.items() for alias in aliases}

def decode_csv(data):
    for encoding in ('utf-8-sig', 'utf-16', 'cp1250', 'cp1251'):
        try:
            result = data.decode(encoding)
            if '\x00' not in result: return result
        except UnicodeError: pass
    raise ValueError('Не могу прочитать кодировку CSV. Экспортируй файл как UTF-8 или XLSX.')

def rows_from_file(data, filename):
    if filename.lower().endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheet = book.active
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            finally: book.close()
        except Exception as exc:
            raise ValueError('Не могу прочитать XLSX. Сохрани выписку заново как XLSX или CSV.') from exc
    elif filename.lower().endswith('.csv'):
        source = decode_csv(data)
        try:
            dialect = csv.Sniffer().sniff(source[:10000], delimiters=',;\t')
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = max((',',';','\t'), key=lambda d: sum(line.count(d) for line in source.splitlines()[:5]))
        try: rows = list(csv.reader(io.StringIO(source, newline=''), delimiter=delimiter))
        except csv.Error as exc: raise ValueError('Ошибка в структуре CSV.') from exc
    else: raise ValueError('Нужен CSV или XLSX. PDF и фото пока не поддерживаются.')
    if len(rows) > 20001: raise ValueError('Слишком много строк; выбери меньший период выписки.')
    if not rows: raise ValueError('Файл пустой.')
    return rows

def parse_money(value):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
    else:
        raw = str(value or '').strip().replace('\u00a0', '').replace('\u202f', '').replace(' ', '').replace('−', '-')
        raw = re.sub(r'(?i)(EUR|USD|GBP|RUB|CHF|€|\$|£|₽)', '', raw)
        negative = raw.startswith('(') and raw.endswith(')')
        raw = raw.strip('()')
        if not re.fullmatch(r'[+-]?\d[\d.,]*', raw): raise ValueError('Некорректная сумма')
        if ',' in raw and '.' in raw:
            decimal_sep = ',' if raw.rfind(',') > raw.rfind('.') else '.'
            raw = raw.replace('.' if decimal_sep == ',' else ',', '')
            raw = raw.replace(decimal_sep, '.')
        elif ',' in raw:
            parts = raw.split(',')
            raw = raw.replace(',', '' if len(parts[-1]) == 3 else '.')
        elif raw.count('.') == 1 and len(raw.split('.')[-1]) == 3:
            raw = raw.replace('.', '')
        number = Decimal(raw) * (-1 if negative else 1)
    if not number.is_finite() or number * 100 != int(number * 100): raise ValueError('Сумма должна иметь не больше двух знаков после запятой')
    return int(number * 100)

def parse_day(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    raw = str(value or '').strip()
    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%d/%m/%y'):
        try: return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError: pass
    raise ValueError('Неизвестный формат даты (или неоднозначное ММ/ДД/ГГГГ)')

def parse_statement(data, filename, default_currency):
    rows = rows_from_file(data, filename)
    choices = []
    for index, row in enumerate(rows[:20]):
        fields = {LOOKUP.get(normal(value)): col for col, value in enumerate(row) if LOOKUP.get(normal(value))}
        if 'day' in fields and ('amount' in fields or 'debit' in fields or 'credit' in fields):
            choices.append((len(fields), index, fields))
    if not choices:
        header = ' | '.join(str(v) for v in rows[0][:12])[:250]
        raise ValueError('Не могу распознать колонки даты и суммы. Пришли названия столбцов выписки (без личных данных), и мы добавим этот формат. Первые поля: ' + header)
    _, header_index, fields = max(choices, key=lambda x: (x[0], -x[1]))
    if 'amount' in fields and ('debit' in fields or 'credit' in fields):
        raise ValueError('В выписке одновременно есть общая сумма и отдельные дебет/кредит. Нужен образец заголовков, чтобы выбрать правильные столбцы.')
    if 'amount' in fields and 'direction' not in fields:
        # A signed amount is clear when at least one row is negative; all-positive files are ambiguous.
        pass
    if not re.fullmatch('[A-Z]{3}', default_currency or ''): raise ValueError('Выбери валюту через /currency перед импортом.')
    result, skipped, occurrences, signed_values = [], 0, {}, []
    for row_number, row in enumerate(rows[header_index + 1:], header_index + 2):
        if not any(str(v or '').strip() for v in row): continue
        def cell(key):
            col = fields.get(key)
            return row[col] if col is not None and col < len(row) else None
        if not cell('day') and not cell('amount') and not cell('debit') and not cell('credit'):
            skipped += 1
            continue
        state = normal(cell('status') or '')
        if state in ('pending', 'reverted', 'rejected', 'cancelled', 'v teku', 'zavrnjeno', 'ожидает', 'отменено'):
            skipped += 1
            continue
        try:
            day = parse_day(cell('day'))
            if 'amount' in fields:
                if cell('amount') in (None, ''): raise ValueError('Сумма отсутствует')
                cents = parse_money(cell('amount'))
                signed_values.append(cents)
                direction = normal(cell('direction') or '')
                if direction in ('debit', 'expense', 'payment', 'withdrawal', 'odliv', 'breme', 'расход', 'списание'):
                    cents = -abs(cents)
                elif direction in ('credit', 'income', 'deposit', 'priliv', 'dobro', 'доход', 'поступление'):
                    cents = abs(cents)
            else:
                debit = parse_money(cell('debit')) if str(cell('debit') or '').strip() else 0
                credit = parse_money(cell('credit')) if str(cell('credit') or '').strip() else 0
                if debit and credit: raise ValueError('Одновременно заполнены расход и приход')
                cents = -abs(debit) if debit else abs(credit)
            if not cents:
                skipped += 1
                continue
            currency = str(cell('currency') or default_currency).strip().upper()
            currency = CURRENCIES.get(currency, currency)
            if not re.fullmatch('[A-Z]{3}', currency): raise ValueError('Неизвестная валюта')
        except (ValueError, InvalidOperation) as exc:
            raise ValueError(f'Строка {row_number}: {exc}. Исправь выписку и загрузи снова; ничего не записано.') from exc
        description = str(cell('description') or cell('direction') or 'Операция по выписке').strip()[:240]
        account = str(cell('account') or '').strip()
        identity = json.dumps([day, cents, currency, description, account, str(cell('balance') or '')], ensure_ascii=False)
        occurrences[identity] = occurrences.get(identity, 0) + 1
        external = hashlib.sha256(f'{identity}:{occurrences[identity]}'.encode()).hexdigest()
        account_id = 'statement:' + hashlib.sha256((account or json.dumps(sorted(fields))).encode()).hexdigest()[:20]
        result.append((external, account_id, day, cents, currency, description, None))
    if not result: raise ValueError('В выписке нет завершённых операций с суммой.')
    if 'amount' in fields and 'direction' not in fields and signed_values and all(c >= 0 for c in signed_values):
        raise ValueError('Все суммы положительные, а колонки расход/приход нет. Не могу определить, где покупки; пришли названия столбцов для настройки формата. Ничего не записано.')
    return result, skipped, 'валюта из столбца' if 'currency' in fields else f'валюта по умолчанию {default_currency}'
